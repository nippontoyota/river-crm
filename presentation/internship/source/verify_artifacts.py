from pathlib import Path
from collections import Counter
import json, re, subprocess, xml.etree.ElementTree as ET
from pptx import Presentation

P=Path(__file__).resolve().parents[1]
data=json.loads((P/'source/slide_content.json').read_text())
deck=Presentation(P/'Nippon_Toyota_Internship_Presentation.pptx')
assert len(deck.slides)==len(data)==46
for i,(slide,d) in enumerate(zip(deck.slides,data),1):
 assert d['notes'] in slide.notes_slide.notes_text_frame.text, ('notes',i)
 for sh in slide.shapes:
  assert sh.left>=0 and sh.top>=0,('negative position',i)
  assert sh.left+sh.width<=deck.slide_width+10000,('horizontal overflow',i)
  assert sh.top+sh.height<=deck.slide_height+10000,('vertical overflow',i)
pdf=P/'Nippon_Toyota_Internship_Presentation.pdf'
pages=subprocess.check_output(['pdftotext',str(pdf),'-']).decode().split('\f')
assert len(pages)-1==46
normalise=lambda t:re.findall(r'[a-z0-9]+',t.casefold().replace('-','').replace('\u00ad',''))
issues=[]
for i,(d,p) in enumerate(zip(data,pages),1):
 missing=Counter(normalise(' '.join(d['text'])))-Counter(normalise(p))
 if missing:issues.append((i,dict(missing)))
assert not issues,issues
root=ET.fromstring(subprocess.check_output(['pdftotext','-bbox-layout',str(pdf),'-']))
ns={'h':'http://www.w3.org/1999/xhtml'}
for i,p in enumerate(root.findall('.//h:page',ns),1):
 for w in p.findall('.//h:word',ns):
  assert 0<=float(w.attrib['xMin'])<=float(w.attrib['xMax'])<=float(p.attrib['width']),('PDF x',i,w.text)
  assert 0<=float(w.attrib['yMin'])<=float(w.attrib['yMax'])<=float(p.attrib['height']),('PDF y',i,w.text)
print('PASS: 46 slides/pages; all embedded notes; shape boundaries; all slide text retained in PDF; PDF text inside page bounds.')
