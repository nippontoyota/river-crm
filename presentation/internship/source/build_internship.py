"""Editable internship presentation; run with python-pptx, reportlab and Pillow."""
from pathlib import Path
import json
from xml.sax.saxutils import escape
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor

P=Path(__file__).resolve().parents[1]
prs=Presentation(); prs.slide_width=Inches(13.333333); prs.slide_height=Inches(7.5)
prs.core_properties.title='Nippon Toyota | Internship Presentation | Mobility CRM'
prs.core_properties.subject='Company orientation, weekly learning, CRM workflows and technical architecture'
prs.core_properties.author='Internship presentation'
C={'ink':'172028','red':'C8202F','paper':'F5F3EF','white':'FFFFFF','muted':'59636A','line':'DADDDC','tint':'F8E6E6','green':'216B59','pale':'E5EFEB','blue':'285D80','dark':'26333D'}
slides=[]; current=None
SOURCES={
 'company':'https://www.nippon-toyota.com/about-us.html',
 'leadership':'https://www.nippon-toyota.com/cms/about-us/dealer-principal-v1.txt',
 'mission':'https://www.nippon-toyota.com/cms/about-us/dealer-mission-v1.txt',
 'location':'https://nippon-toyota.com/contact-co01b.html',
 'history':'https://www.linkedin.com/company/nippon-toyota-pvt-ltd',
 'used':'https://www.nippon-toyota.com/used-cars.html',
 'photo1':'https://commons.wikimedia.org/wiki/File:Kalamassery_Toyota_Office.JPG',
 'photo2':'https://www.flickr.com/photos/sreejithmsivadasan/5974734010',
}
def rgb(c):return RGBColor.from_string(C.get(c,c))
def box(s,x,y,w,h,fill='white',border=None):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
 if border:sh.line.color.rgb=rgb(border)
 else:sh.line.fill.background()
 sh._element.spPr.append(OxmlElement('a:effectLst'))
 return sh
def text(s,x,y,w,h,value,size=20,color='ink',bold=False,center=False):
 current['text'].append(str(value))
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
 tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(str(value).split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Lato';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(7)
  if center:p.alignment=PP_ALIGN.CENTER
 return sh
def line(s,x1,y1,x2,y2,color='red',width=2,arrow=False,dash=False):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(color);sh.line.width=Pt(width)
 if dash:
  from pptx.enum.dml import MSO_LINE_DASH_STYLE
  sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
 if arrow:
  a=OxmlElement('a:tailEnd');a.set('type','triangle');sh._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln').append(a)
 return sh
def new(title,section,sub='',notes='',refs=(),dark=False):
 global current
 s=prs.slides.add_slide(prs.slide_layouts[6]);current={'number':len(slides)+1,'title':title,'section':section,'notes':notes,'refs':list(refs),'text':[]};slides.append(current)
 s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('ink' if dark else 'paper')
 box(s,.0,.0,.16,7.5,'red');text(s,.6,.37,11,.3,section.upper(),11,'red',True)
 text(s,.6,.95,12.15,.62,title,31,'white' if dark else 'ink',True)
 if sub:text(s,.62,1.72,12,.58,sub,16,'line' if dark else 'muted')
 line(s,.6,6.91,12.72,6.91,'dark' if dark else 'line',1)
 text(s,.62,7.12,10,.2,'NIPPON TOYOTA  /  KALAMASSERY  /  INTERNSHIP',9,'line' if dark else 'muted')
 text(s,12.02,7.05,.65,.3,f'{len(slides):02d}',13,'red',True,True)
 return s
def band(s,value,color='red'):
 box(s,.62,6.24,12.08,.46,'tint' if color=='red' else 'pale');text(s,.8,6.34,11.74,.25,value,12,color,True)
def card(s,x,y,w,h,title,body,accent='red',size=19):
 box(s,x,y,w,h,'white');box(s,x,y,.045,h,accent)
 text(s,x+.2,y+.2,w-.4,.65,title,20,accent,True)
 offset=.72 if h<2 else .96
 text(s,x+.2,y+offset,w-.4,h-offset-.08,body,size if h>=2 else 15,'ink')
def cards(title,section,sub,items,notes='',refs=(),take=''):
 s=new(title,section,sub,notes,refs)
 n=len(items);gap=.22;w=(12.08-gap*(n-1))/n
 for i,(h,b) in enumerate(items):card(s,.62+i*(w+gap),2.57,w,3.29,h,b,size=19 if n<3 else 18)
 if take:band(s,take)
 return s
def table(title,section,sub,rows,widths,notes='',refs=(),size=16):
 s=new(title,section,sub,notes,refs);tb=s.shapes.add_table(len(rows),len(rows[0]),Inches(.62),Inches(2.49),Inches(12.08),Inches(3.52)).table
 for i,w in enumerate(widths):tb.columns[i].width=Inches(w)
 for r,row in enumerate(rows):
  for c,v in enumerate(row):
   current['text'].append(v);ce=tb.cell(r,c);ce.text=v;ce.fill.solid();ce.fill.fore_color.rgb=rgb('ink' if r==0 else ('white' if r%2 else 'EAEDEB'));ce.vertical_anchor=MSO_ANCHOR.MIDDLE
   ce.margin_left=Inches(.13);ce.margin_right=Inches(.1);ce.margin_top=ce.margin_bottom=Inches(.06)
   for p in ce.text_frame.paragraphs:p.font.name='Lato';p.font.size=Pt(size);p.font.bold=r==0 or c==0;p.font.color.rgb=rgb('white' if r==0 else 'ink')
 return s
def node(s,x,y,w,title,body='',h=1.05,color='red'):
 box(s,x,y,w,h,'white','line');text(s,x+.14,y+.13,w-.28,.42,title,19,color,True,True)
 if body:text(s,x+.14,y+.6,w-.28,h-.63,body,14,'muted',center=True)
def photo(s,path,x,y,w,h):
 iw,ih=Image.open(path).size;scale=min(w/iw,h/ih);ww,hh=iw*scale,ih*scale
 s.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))

# 01–08: Company and workplace.
s=new('Internship at Nippon Toyota','Internship presentation','Kalamassery showroom, Kochi, Kerala',
 'Introduce yourself, your course and your internship dates. Explain that the mobility CRM is one project undertaken during the internship. The workspace calls it Incheon Mobility CRM and contains River scooter configuration; Nippon Toyota is the internship host supplied by the presenter. Do not imply this is Toyota’s official enterprise CRM. Personal details and the eight-week allocation remain editable until confirmed.',dark=True)
text(s,.68,2.75,6.2,1.35,'Mobility CRM\nDevelopment & Learning',37,'white',True)
text(s,.7,4.47,5.9,1.17,'[Your name]  |  [Register number]\n[Course and college]\n[Internship dates]  |  [Mentor]',17,'line')
photo(s,P/'assets/kalamassery-showroom-2012.jpg',7.25,2.55,5.38,3.66)
text(s,7.3,6.39,5.2,.24,'Kalamassery, 2012 • Ranjithsiji • CC BY-SA 3.0',10,'line')

cards('Presentation route','Overview','Company context → weekly progress → technical work → learning',[
 ('Workplace','Company profile\nKalamassery showroom\nLeadership and departments'),('Weekly journey','Week 1: induction\nWeek 2: R&D\nWeeks 3–8: project stages'),('Project and learning','Workflows and architecture\nDatabase, APIs and security\nTesting, challenges and reflection')],
 'The deck has a main internship narrative followed by a technical appendix for viva questions. Week 1 and Week 2 themes come directly from your request. The later weekly allocation is a proposed reconstruction, not an attendance record. Explain that the project discussion covers the current repository and that your exact personal contribution should be confirmed with your internship diary.',take='Timeline basis: eight-week draft; replace with your actual internship dates and diary entries.')

cards('Nippon Toyota: company profile','Company','Internship host: Nippon Toyota | Legal entity listed: Nippon Motor Corporation Pvt. Ltd.',[
 ('Automotive retail','Toyota dealership operations\nVehicle sales, servicing and spares\nCustomer relationships across the ownership cycle'),('History and presence','Company LinkedIn profile lists establishment on 27 January 2000\nLocations across Kerala'),('Customer support','Sales and after-sales services\nFinance and insurance information\nUsed-car offering through Nippon U Trust')],
 'Distinguish Toyota, the vehicle brand, from Nippon Toyota, the dealership business hosting the internship. The dealer contact page identifies Nippon Motor Corporation Pvt. Ltd. The company’s own LinkedIn profile gives 27 January 2000 as its establishment date. The official dealership site lists locations across Kerala and the used-car page lists Nippon U Trust at Kalamassery. Avoid unverified employee counts, turnover or market rankings.',[SOURCES['location'],SOURCES['history'],SOURCES['company'],SOURCES['used']])

cards('Customer service shapes the work','Company','The published mission prioritises customer satisfaction and integrated sales, service and parts.',[
 ('Sales','Understand an enquiry\nExplain the vehicle and purchase options\nCoordinate test drive, booking and delivery'),('Service and spares','Receive service requests\nCoordinate maintenance and parts\nRecord issues and follow through'),('CRM relevance','Keep customer context available\nAssign a responsible employee\nRecord the next action and outcome')],
 'Paraphrase the published dealer mission: the company places customer satisfaction first and brings sales, service and service parts together for convenience and efficiency. The department activities are a general dealership process explanation, not a claim that this project implements workshop job cards, inventory, finance underwriting or Toyota integrations. The project focuses on customer enquiries, sales follow-ups and complaints.',[SOURCES['mission']],take='Project boundary: enquiry and complaint operations; vehicle inventory and workshop billing are outside this repository.')

s=new('My workplace: Kalamassery showroom','Workplace','Nippon Towers, NH 47, HMT Junction, Kalamassery P.O., Ernakulam, Cochin 683104',
 'The presenter confirmed this as the internship location. The official location page describes a showroom and service centre at this address. These public photographs show the real Kalamassery facility in 2011 and 2012. They are historical building views, not photographs of the presenter’s desk or the current office interior. If available, replace or supplement them with your own current workplace photographs.',[SOURCES['location'],SOURCES['photo1'],SOURCES['photo2']])
photo(s,P/'assets/kalamassery-showroom-2012.jpg',.62,2.42,6.02,3.58)
photo(s,P/'assets/kalamassery-showroom-2011.jpg',6.91,2.42,5.78,3.58)
text(s,.69,6.07,5.9,.22,'13 Mar 2012 • Ranjithsiji • Wikimedia Commons • CC BY-SA 3.0',9,'muted')
text(s,6.96,6.07,5.65,.22,'22 Jul 2011 • Sreejith Sivadasan • Flickr • All rights reserved',9,'muted')
text(s,.7,6.49,11.8,.24,'Historical exterior photographs. Current office and personal workplace photographs have not been supplied.',11,'red',True)

s=new('Company leadership and departmental hierarchy','Week 1 | company orientation','Published leadership above; illustrative departmental reporting structure below.',
 'The official company site names M. A. M. Babu Moopan as Dealer Principal and Athif Moopan and Naeem Shahul as Directors. The arrangement of branch management and functional departments is an illustrative dealership structure, because an internal organisation chart was not supplied. Do not present the connecting lines as verified reporting relationships. During rehearsal, replace the lower structure with the hierarchy explained during your induction, including your department and reporting mentor.',[SOURCES['leadership']])
node(s,4.03,2.4,5.27,'M. A. M. Babu Moopan','Dealer Principal',1.02)
node(s,.75,3.72,3.42,'Athif Moopan','Director',1.0)
node(s,9.15,3.72,3.42,'Naeem Shahul','Director',1.0)
node(s,4.56,3.72,4.2,'Branch / showroom management','Illustrative reporting layer',1.0)
line(s,6.66,3.42,6.66,3.72,'muted',1.5,dash=True)
for i,(t,b) in enumerate([('Sales / CRM','Sales team · CE · reception'),('Service / spares','Service advisers · parts'),('Finance / HR / IT','Support · project mentor')]):
 x=.75+i*4.18;node(s,x,5.06,3.83,t,b,1.0);line(s,6.66,4.72,x+1.91,5.06,'muted',1.5,dash=True)
band(s,'Confirm the internal reporting lines and your mentor’s position before presenting.')

cards('Internship objectives','Learning goals','Connect business observation with a working software implementation.',[
 ('Understand the workplace','Learn department responsibilities\nObserve enquiry and follow-up practices\nUnderstand escalation and reporting'),('Develop the project','Translate requirements into workflows\nBuild role-specific interfaces and APIs\nPreserve ownership and history'),('Build professional skills','Discuss requirements with stakeholders\nTest changes and explain findings\nDocument progress and technical decisions')],
 'Use these as internship objectives rather than claims about activities you have already completed. Connect your academic knowledge to a business setting: customer data is useful only if employees can use it to make the next decision. Add the objectives agreed with your mentor and the specific responsibilities allocated to you. Explain how you will show evidence through code, tests, workflows and your weekly diary.')

cards('Project context: a mobility CRM','Project introduction','Repository name: Incheon Mobility CRM | Presented as a project undertaken during the Nippon Toyota internship.',[
 ('Business problem','Enquiries need an owner\nFollow-ups need dates and outcomes\nManagers need branch visibility'),('Implemented scope','Lead intake and allocation\nCE-to-PS/SO sales handoff\nComplaints, analytics and staff continuity'),('Naming and ownership','Current code includes River scooter settings\nConfirm the official project name\nConfirm which modules you personally built')],
 'The repository’s README, API schema title and frontend package identify Incheon Mobility CRM. A migration seeds River Indie. The code does not establish a corporate relationship between Incheon Mobility, River and Nippon Toyota. Present this as a mobility CRM project worked on during the internship until the official project/client description is confirmed. Do not claim that all repository features were built by one intern or that the company has adopted the full system.', ['README.md','frontend/package.json','backend/leads/migrations/0018_seed_river_indie.py'])

# 09–17: Weekly allocation, deliberately marked as draft.
weeks=[
 ('Company induction','Introduction to Nippon Toyota; departments, hierarchy and showroom operations','Orientation notes and department map'),
 ('R&D and follow-up','Study the customer journey; clarify requirements and compare technical options','Requirement list and workflow draft'),
 ('Solution design','Define roles, entities, API boundaries and screen structure','Architecture and data model'),
 ('Core implementation','Build intake, allocation, role visibility and employee workspaces','Working enquiry-to-owner flow'),
 ('Sales workflow','Add qualification, follow-ups, booking, retail and test-drive completion','Audited sales workflow'),
 ('Operations modules','Develop complaints, imports, configurable lists and staff lifecycle handling','Operations features and validations'),
 ('Analytics and verification','Review ETBR, branch scope, cache behaviour and regression tests','Validated reporting and defect fixes'),
 ('Delivery and reflection','Review deployment setup, prepare documentation, demo and learning summary','Presentation and handover checklist'),
]
s=new('Weekly internship timeline','Work plan','Eight-week draft allocation | Week 1 and Week 2 themes supplied by the presenter.',
 'This is a proposed way to organise the project into an internship narrative. It is not a statement of actual calendar completion. Replace the week count and allocate each feature using your diary, commits and mentor feedback. If the internship was shorter, combine adjacent implementation stages; if longer, split design, testing or deployment. Week 1 remains induction and hierarchy; Week 2 remains research, requirements and follow-up.')
for i,(h,b,o) in enumerate(weeks):
 x=.7+(i%4)*3.14;y=2.43+(i//4)*1.82
 box(s,x,y,2.94,1.58,'white');text(s,x+.15,y+.12,2.6,.3,f'WEEK {i+1:02d}',12,'red',True);text(s,x+.15,y+.5,2.6,.7,h,18,'ink',True);text(s,x+.15,y+1.28,2.6,.23,'Induction / R&D' if i<2 else 'Proposed allocation',10,'muted')
band(s,'Map these stages to your actual internship dates; no week-by-week completion dates have been invented.')

details=[
 ('Introduction, hierarchy and work culture', [('Activities','Company introduction\nShowroom and service overview\nDepartment and reporting hierarchy'),('Learning focus','Customer service expectations\nEmployee responsibilities\nProfessional communication and data handling'),('Weekly evidence','Induction notes\nDepartment map\n[Add mentor, visit and observation details]')],
 'Your requested first week focuses on the introduction to the company and the hierarchy. Explain the difference between sales, customer relations, reception, service and support teams. Add the names of the people you met and the induction activities you actually attended. A useful experience statement, if accurate, is: “I began to understand how employees coordinate around a single customer enquiry.” Do not imply participation in service work or meetings that did not occur.'),
 ('R&D, requirements and follow-up', [('Activities','Study enquiry-to-sale flow\nIdentify data and reporting needs\nFollow up on unclear requirements'),('Technical exploration','Compare spreadsheet and web workflows\nReview React/Next.js and Django options\nDraft data model and role matrix'),('Weekly evidence','Requirement questions and answers\nWorkflow sketch\n[Add actual discussion or feedback]')],
 'Your requested second week focuses on R&D and follow-up. Describe what was researched, which existing process or form you reviewed, and the questions you took back to your mentor. Follow-up here means clarifying requirements as well as understanding customer follow-up operations. Suggested reflection, if accurate: “I learned to ask who performs each action and what happens when the customer does not respond.” Technology comparisons are proposed talking points unless your diary confirms you made them.'),
 ('Designing the solution', [('Activities','Break requirements into modules\nMap roles and visibility\nSketch screens and API interactions'),('Technical work','User and Lead relationships\nSeparate lead status from sales outcome\nPlan validation and audit events'),('Weekly evidence','Architecture diagram\nEntity relationship diagram\nAPI and screen inventory')],
 'Proposed Week 3. Describe how the data model supports the workflow: a lead retains qualification, multiple calls, multiple follow-ups and audit events. The application distinguishes the customer executive from the sales employee. Discuss why sales outcome is separate from the initial contact status. Confirm whether these were design decisions you made or design elements you studied. Cite the models and routes as technical evidence, not as proof that they were completed in this exact week.'),
 ('Building intake and allocation', [('Activities','Implement role-specific forms\nCapture digital and walk-in enquiries\nRoute records to eligible employees'),('Technical work','Django serializers and viewsets\nTyped frontend API calls\nOwnership and branch checks'),('Weekly evidence','Lead creation and assignment flow\nValidation examples\nRole-based queue review')],
 'Proposed Week 4. Explain the three important entry patterns: admin/CE intake, receptionist walk-in capture and SO self-generated enquiries. Admins can allocate fresh leads to customer executives. Qualification then assigns the sales employee. The stored field assigned_so refers to the CE allocation, while assigned_ps refers to PS/SO ownership. This naming detail matters when explaining the implementation. Use a fictional record when demonstrating; describe only the screens or endpoints you actually implemented.'),
 ('Implementing the sales journey', [('Activities','Record qualification and call outcomes\nSchedule follow-ups\nTrack booking, retail and test drives'),('Technical work','State validation and transactions\nCall logs and before/after audit\nIdempotent test-drive completion'),('Weekly evidence','CE-to-PS/SO handoff\nPending and lost scenarios\nBooked-to-retailed walkthrough')],
 'Proposed Week 5. Follow one lead through qualification, pending sales discussion, booking and retail. Explain that a requested test drive is different from a completed one; the current completion action stores an audited timestamp once. Call outcomes and next follow-up rules keep queues meaningful. A booking does not automatically prove delivery or payment integration. Relate the code to the module you actually worked on and place later additions into their correct internship week.'),
 ('Adding operational support', [('Activities','Complaint intake and resolution\nCSV/XLSX import review\nLists and employee lifecycle controls'),('Technical work','Type/subtype and activity validation\nDuplicate-phone staging\nSafe reassignment and reminder holds'),('Weekly evidence','Complaint workflow\nImport summary and commit\nOffboarding impact preview')],
 'Proposed Week 6. Explain how operational exceptions influence software design. Complaints have a separate lifecycle and permissions. Imports need a review stage because a successful upload can still contain invalid rows or duplicates. Staff departure must preserve history while moving active work to another owner or queue. Do not imply that the application deletes historical customer data when an employee account is offboarded. Match the weekly account to your actual contribution and progress.'),
 ('Analytics, testing and refinement', [('Activities','Review enquiry-stage metrics\nCheck branch and role scoping\nRun regression and frontend checks'),('Technical work','Distinct ETBR aggregates\nFilter-aware cache keys\nPermission and workflow tests'),('Weekly evidence','Test results and bug notes\nReporting definitions\nVerified edge cases')],
 'Proposed Week 7. Explain the tests that matter to users: can someone see another employee’s lead, can a manager see another branch, does a repeated test-drive click double count, and do configuration changes preserve historical records? The current verification run belongs to preparation of this presentation, not necessarily to your original Week 7. State that distinction. No production conversion rate or response-time improvement has been measured for the presentation.'),
 ('Delivery, documentation and review', [('Activities','Review environment configuration\nPrepare demo and technical notes\nSummarise learning and remaining work'),('Technical work','Frontend and API deployment boundaries\nDatabase migration and health checks\nWorker/scheduler configuration review'),('Weekly evidence','Demo script and handover checklist\nArchitecture and API notes\n[Add mentor feedback and final review]')],
 'Proposed Week 8. The repository includes a Render service definition and Vercel frontend configuration. Describe configuration and deployment work you actually performed. The checked-in Render setup enables eager Celery execution and does not define a worker or scheduler, so do not claim background reminders are live. For handover, explain the need to apply migrations, configure secrets, verify each role and document unresolved items. Add real mentor feedback rather than an invented endorsement.'),
]
for i,(h,items,note) in enumerate(details):
 cards(f'Week {i+1:02d}: {h}','Weekly learning','Presenter-supplied theme; details to personalise.' if i<2 else 'Proposed allocation: verify against your internship diary.',items,note,take='Output: '+weeks[i][2])

# 18 onwards: Detailed technical walkthrough.
table('Requirements translated into features','Analysis','Business needs mapped to implementation and acceptance evidence.',[
 ['Need','Implemented response','Acceptance example'],
 ['Clear ownership','CE and PS/SO assignment','Assigned employee sees the same lead'],
 ['Consistent follow-up','Outcome, date, remarks and history','Invalid follow-up date is rejected'],
 ['Controlled access','Role and branch-scoped API queries','Other branch records stay outside the result'],
 ['Reliable intake','Configured lists and import staging','Invalid values and duplicate rows are identified'],
 ['Management review','ETBR and role-specific analytics','Same scope and dates give comparable totals'],
 ['Staff continuity','Impact preview and reassignment','Active work remains actionable after disable'],
],[2.35,4.6,5.13],
 'Frame the problem as project requirements. The repository does not prove the company previously lost enquiries or used a particular spreadsheet, so avoid inventing an as-is failure story. For each requirement, identify a visible behaviour or an API response that can confirm it. These acceptance examples are stronger than unsupported claims of improved sales. Link personal contribution to the requirements you actually handled.', ['backend/leads/views.py','backend/accounts/offboarding.py','backend/uploads/tasks.py','backend/analytics/test_etbr.py'],15)

s=new('End-to-end customer enquiry workflow','Workflow','A lead keeps its identity while responsibility and sales progress change.',
 'Read the main path from left to right. Admin or CE capture and bulk imports feed the enquiry pool. A CE contacts the customer and qualifies an interested enquiry, retaining the original CE relationship while handing sales work to a PS/SO. The PS/SO records follow-up, booking and retail. Receptionist walk-ins and SO-generated enquiries can enter the sales route directly. A complaint is a separate ticket and is not an automatic transition from a lost lead.', ['backend/leads/views.py','backend/leads/serializers.py','frontend/src/features/leads/sales-workspace.tsx'])
for i,(h,b) in enumerate([('Capture','Admin / CE / import'),('CE contact','Assign and qualify'),('PS/SO follow-up','Discuss / test drive'),('Sale outcome','Booked → Retailed')]):
 x=.7+i*3.16;node(s,x,2.78,2.85,h,b,1.27)
 if i<3:line(s,x+2.85,3.41,x+3.13,3.41,arrow=True)
node(s,.7,4.73,3.7,'Walk-in / SO enquiry','Direct sales entry',1.05,'blue')
line(s,4.4,5.25,8.44,5.25,'blue',2,arrow=True);line(s,8.44,5.25,8.44,4.05,'blue',2,arrow=True)
node(s,4.85,4.05,2.87,'Pending or Lost','Reason and history',1.0,'muted')
band(s,'Managers review authorised branch records; complaints follow their own ticket workflow.')

table('Roles and access boundaries','Access control','Six application roles; the CE role still uses the stored code CRE.',[
 ['Role','Main responsibility','Visibility / restriction'],
 ['Admin','Configure, allocate and oversee','Broad CRM access; complaints read-only'],
 ['CE (CRE)','Contact, qualify and log issues','Assigned leads; own logged complaints'],
 ['PS/SO (SO)','Sales follow-up, booking and retail','Assigned sales leads and own intake'],
 ['Sales Manager','Review sales performance','Configured branch; management views'],
 ['Receptionist','Capture walk-ins and complaints','Own capture dashboard and tickets'],
 ['Complaints team','Investigate and resolve issues','Complaint handling and complaint analytics'],
],[2.1,4.5,5.48],
 'Do not confuse the company organisation chart with the application permission matrix. The browser chooses the workspace, while the backend enforces authorisation. CE is the display label for the legacy CRE role. Lead queries use assigned_so for CE ownership and assigned_ps for PS/SO. A sales manager without a configured location receives no branch leads. The complaint permission class gives admins list/retrieve access, CE and receptionists create/read access, and the complaints department update and notes access.', ['backend/accounts/models.py','backend/leads/views.py','backend/complaints/permissions.py'],15)

table('Technology stack and responsibilities','Architecture','Versions and capabilities described from the checked-in project configuration.',[
 ['Layer','Technology','Responsibility'],
 ['Interface','Next.js, React, TypeScript, CSS','Role workspaces, forms, dashboards and routing'],
 ['API','Python, Django 5.2 range, DRF','Validation, business actions and visibility'],
 ['Identity','SimpleJWT + bcrypt-SHA256','Cookie authentication and password hashing'],
 ['Data','Django ORM; environment-selected SQL DB','Relations, constraints, transactions and migrations'],
 ['Files / tasks','openpyxl; Supabase storage; Celery','Spreadsheet parsing, file storage and tasks'],
 ['Operations','Redis cache; Gunicorn/Uvicorn; OpenAPI','Analytics caching, API serving and documentation'],
],[1.75,4.83,5.5],
 'The frontend manifest uses latest for Next.js, React and TypeScript, so this slide does not invent pinned versions. The backend requirements constrain Django to the 5.2 series and include Django REST Framework, SimpleJWT, Celery, Redis integration, openpyxl, psycopg and Supabase. SQLite is the default local database; DATABASE_URL chooses the deployed database. PostgreSQL is supported by the installed driver, but the live database provider was not inspected. External integrations shown as lead sources are labels, not proof of automatic ingestion.', ['frontend/package.json','backend/requirements.txt','backend/config/settings.py'],15)

s=new('System architecture','Architecture','Browser → Django REST API → data and supporting services',
 'The frontend communicates with Django rather than calling the database directly. Django owns authentication, request validation, record visibility and business transactions. File uploads use Supabase storage when credentials are configured, with a local storage fallback. Analytics can use Redis or local memory. Celery tasks exist, but deployment mode determines whether they run inline or on a worker. Dashed connections identify optional infrastructure, not verified live services.', ['frontend/src/lib/crm.ts','backend/config/settings.py','backend/uploads/storage.py','render.yaml'])
node(s,.7,3.1,2.75,'Browser UI','Next.js / React / TS',1.15,'blue')
node(s,4.27,3.1,4.37,'Django REST API','Auth · serializers · roles · workflows',1.15)
node(s,9.45,3.1,3.12,'SQL database','ORM · audit · transactions',1.15,'green')
line(s,3.45,3.67,4.27,3.67,'blue',2,arrow=True);line(s,8.64,3.67,9.45,3.67,'green',2,arrow=True)
node(s,4.27,4.93,3.8,'File storage','Supabase / local fallback',1.02,'blue')
node(s,8.6,4.93,3.97,'Cache / task services','Redis / Celery when configured',1.02,'muted')
line(s,6.17,4.25,6.17,4.93,'blue',2,arrow=True)
line(s,7.73,4.25,10.58,4.93,'muted',2,arrow=True,dash=True)
text(s,4.37,2.48,4.16,.28,'HTTPS JSON + authenticated cookies',13,'muted',center=True)
band(s,'The checked-in Render service uses eager Celery tasks; a separate worker and scheduler are not declared.')

cards('Frontend structure and API communication','Implementation','Feature-based UI with a shared API helper and role-specific routes.',[
 ('Screens and components','App Router role groups\nSales workspace and lead desk\nShared shell, forms and ETBR tiles'),('Typed API layer','Request/response types in crm.ts\nSnake-case API mapping to UI fields\nCredentials and CSRF headers'),('Error and session handling','Field errors shown to users\nRefresh on an unauthorised response\nOne retry after successful refresh')],
 'The frontend uses route groups for admin, sales, manager and receptionist screens. Feature directories contain the lead desk, sales workspace, complaints, analytics and team pages. src/lib/crm.ts defines request types and response mappings. The api helper includes credentials, adds CSRF headers for unsafe methods and retries once after a successful refresh response. sessionStorage holds cached user metadata, not the JWT. Backend checks remain necessary even if a link or button is absent from a user’s workspace.', ['frontend/src/lib/crm.ts','frontend/src/components/app-shell.tsx','frontend/src/app','frontend/src/features'])

cards('Backend modules and request processing','Implementation','Django apps separate customer operations from shared platform concerns.',[
 ('Domain modules','accounts: roles and staff lifecycle\nleads: ownership and sales actions\ncomplaints: ticket workflow'),('Supporting modules','uploads: staging and import\nanalytics: scoped aggregates\nnotifications: due follow-up records'),('Request path','URL/router selects the action\nSerializer validates the payload\nAction checks permissions and saves')],
 'Django REST Framework viewsets provide resource endpoints and dedicated actions for behaviours such as qualification, assignment and test-drive completion. Serializers validate configured choices and outcome-specific fields. The action then checks ownership and performs related changes in a transaction. Not every model operation has the same rules, so prefer the dedicated workflow endpoints when explaining business behaviour. drf-spectacular publishes the OpenAPI schema and Swagger UI at /api/schema/ and /api/docs/.', ['backend/config/urls.py','backend/leads/urls.py','backend/leads/serializers.py','backend/leads/views.py'])

s=new('Core database relationships','Data model','Simplified entity relationship diagram; connectors show relationships and cardinality.',
 'User has two lead ownership relationships: assigned_so is the CE and assigned_ps is the sales employee. LeadQualification is optional one-to-one with Lead. A lead can have many call logs, follow-ups and audit events. Complaints can optionally reference a lead and have their own notes. This slide simplifies fields and omits additional actor foreign keys for readability. The appendix lists the other operational entities. Protected references preserve history when an account or lead is involved in audit-related records.', ['backend/accounts/models.py','backend/leads/models.py','backend/complaints/models.py'])
node(s,.75,2.75,3.35,'User','id · email · role · location',1.1,'blue')
node(s,5.0,2.75,3.4,'Lead','id · uid · status · sales_outcome',1.1)
node(s,9.25,2.75,3.3,'LeadQualification','lead_id (unique) · finance · notes',1.1,'green')
line(s,4.1,3.3,5,3.3,'blue',2,arrow=True);text(s,4.14,2.72,.8,.3,'1 → N',11,'blue',center=True)
line(s,8.4,3.3,9.25,3.3,'green',2,arrow=True);text(s,8.43,2.72,.82,.3,'1 → 0..1',10,'green',center=True)
for i,(h,b) in enumerate([('CallLog','lead_id · outcome · remarks'),('FollowUp','lead_id · owner · scheduled_for'),('LeadAudit','lead_id · actor · before / after')]):
 x=.75+i*4.21;node(s,x,4.91,3.8,h,b,1.05);line(s,6.7,3.85,x+1.9,4.91,'muted',1.3,arrow=True);text(s,x+1.34,4.56,1.12,.25,'1 → many',10,'muted',center=True)
band(s,'Each lead can also have many complaints; each complaint can have many complaint notes.')

table('Lead fields and business meaning','Data model','Two ownership fields and separate progress fields preserve the handoff context.',[
 ['Field / relationship','Purpose','Implementation detail'],
 ['assigned_so / assigned_ps','CE owner / PS-SO owner','Both point to accounts.User; history protected'],
 ['status / sales_outcome','Contact stage / sales progress','Separate enums; category stores Hot/Warm/Cold'],
 ['source / activity / sub_activity','Origin and marketing context','Sub-activity belongs to a configured activity'],
 ['branch / rto / enquiry_date','Location and reporting context','RTO choices; date validation and filters'],
 ['test_drive_completed_at','Recorded completion evidence','One audited timestamp; repeated action is safe'],
 ['deleted_at / reassignment flags','Record lifecycle and work routing','Soft deletion and explicit reassignment queues'],
],[3.16,3.85,5.07],
 'The Lead model has a numeric primary key and a public UUID. Phone is indexed but not unique, so do not claim database-level duplicate prevention for manual intake. Qualification is a linked entity; historical activity labels remain stored on leads when configurable choices are retired. Indexes combine CE or PS ownership with status, and include source with creation time. The current data model stores timestamps and status snapshots, but does not implement a full event-sourced analytics warehouse.', ['backend/leads/models.py','backend/leads/serializers.py'],15)

s=new('CE qualification and sales-state workflow','Workflow','Qualification starts sales work; booking and retail are separate outcomes.',
 'A fresh lead may move through response-related states such as RNR, switched off, callback or pending. For the normal CE route, qualification captures the buying context and assigns the PS/SO. Lost is a reasoned outcome with remarks. The PS/SO has separate sales outcomes: pending, booked, retailed and lost. This diagram shows the main business path rather than every server transition. A completed test drive is a parallel timestamp and does not change lead status, qualification or the open follow-up.', ['backend/leads/views.py','backend/leads/serializers.py','backend/leads/metrics.py'])
node(s,.7,2.7,2.25,'Fresh','New enquiry',1.07)
node(s,4.0,2.7,3.42,'Qualified','CE context + named PS/SO',1.07,'green')
node(s,8.34,2.7,4.24,'Sales progress','Pending → Booked → Retailed',1.07,'green')
line(s,2.95,3.22,4,3.22,arrow=True);line(s,7.42,3.22,8.34,3.22,'green',2,arrow=True)
node(s,.7,4.74,3.72,'Contact pending','RNR / switched off / callback',1.12,'blue')
line(s,1.84,3.77,1.84,4.74,'blue',2,arrow=True)
node(s,5.02,4.74,2.87,'Lost','Reason + remarks',1.12)
line(s,5.7,3.77,6.45,4.74,'red',2,arrow=True)
node(s,8.48,4.74,4.1,'Test drive completed','Audited event alongside the pipeline',1.12,'blue')
band(s,'Requested test drive ≠ completed test drive. A completed drive does not itself mark the lead booked.')

s=new('Follow-up loop and reminder conditions','Workflow','A next action needs an owner, a future date and an outcome when the call is made.',
 'For the CE update path, the serializer rejects appointments in the past and more than three days ahead. Pending, callback and walk-in outcomes require a next appointment. Other paths have their own validations; do not claim identical validation across every endpoint. When processing due reminders, the task excludes resolved, already notified, held, deleted-lead and inactive-owner items. It records a notification and marks notified_at. A periodic schedule exists in settings, but delivery still needs a running scheduler or equivalent trigger.', ['backend/leads/serializers.py','backend/leads/views.py','backend/notifications/tasks.py','backend/config/settings.py'])
for i,(h,b) in enumerate([('Record outcome','Call result + remarks'),('Set next action','Owner + scheduled date'),('Review due queue','Due / overdue follow-up'),('Call and update','Resolve / replace next action')]):
 x=.7+i*3.16;node(s,x,2.9,2.85,h,b,1.24)
 if i<3:line(s,x+2.85,3.5,x+3.14,3.5,arrow=True)
line(s,11.58,4.14,11.58,4.65,'blue',2);line(s,11.58,4.65,2.13,4.65,'blue',2);line(s,2.13,4.65,2.13,4.14,'blue',2,arrow=True)
text(s,.92,5.17,11.4,.6,'Reminder task: unresolved + due + active owner + not held + not already notified',18,'blue',True,True)
band(s,'Scheduler setting: every 900 seconds. Live reminder execution has not been verified.')

s=new('Authentication request sequence','Security','Cookie-based JWT authentication, CSRF checks and one refresh retry in the frontend.',
 'The login endpoint validates credentials and returns user metadata while setting HttpOnly access and refresh cookies. The access lifetime is 15 minutes and refresh lifetime is 7 days. In non-debug mode cookies are Secure and SameSite=None. For cookie-authenticated unsafe requests, CookieJWTAuthentication invokes a CSRF check. Refresh blacklists the old refresh token and issues a new one. These controls do not constitute a full security audit; login and refresh are custom endpoints and should be considered separately during hardening.', ['backend/accounts/authentication.py','backend/accounts/views.py','backend/config/settings.py','frontend/src/lib/crm.ts'])
xs=[1.6,6.6,11.6]
for x,h in zip(xs,['Browser','Django API','Database']):
 text(s,x-1,2.4,2,.4,h,20,'red',True,True);line(s,x,2.95,x,5.85,'line',1.5)
for y,a,b,t in [(3.17,0,1,'1  POST login: credentials'),(3.8,1,2,'2  Load account; API verifies password'),(4.43,1,0,'3  Set HttpOnly cookies + return user'),(5.06,0,1,'4  API request: cookie + CSRF header'),(5.69,1,0,'5  Authorised JSON response')]:
 line(s,xs[a],y,xs[b],y,'blue' if a==0 else 'red',2,arrow=True)
 text(s,min(xs[a],xs[b])+.14,y-.38,abs(xs[b]-xs[a])-.25,.3,t,13,'ink',center=True)
band(s,'Access: 15 minutes | Refresh: 7 days | Password hashing: bcrypt-SHA256 | Backend role checks remain essential.')

cards('Security and data integrity controls','Security','Existing protections and the reasons they matter.',[
 ('Trust boundaries','Authenticated API by default\nRole and branch checks on queries\nConfigured values validated on input'),('Consistent updates','transaction.atomic for related writes\nselect_for_update on sensitive actions\nAudit actor, event and before/after data'),('Operational controls','Secrets read from environment\nRequest throttling configured\nProtected history and lifecycle records')],
 'Explain transaction atomicity using a lead update: the business status and its related history should either save together or fail together. Row locking helps concurrent operations coordinate where the action uses it. The API config sets anonymous and authenticated throttle rates, but this is not a guarantee against every abuse scenario. CORS is broad for Vercel preview origins in the current configuration, and process-local throttling and live settings deserve deployment review. Never show .env contents or production credentials in the presentation.', ['backend/config/settings.py','backend/leads/views.py','backend/accounts/offboarding.py'],take='Security review scope: code-level controls and workflow tests; no penetration-test or compliance certification claim.')

s=new('CSV / XLSX import pipeline','Data ingestion','Review staged rows before adding them to the active CRM.',
 'An admin uploads a CSV or XLSX file through a multipart request. The API checks extension and a 10 MB size limit, stores the file and calls the parsing task. The parser normalises Indian phone numbers and validates name, phone, source, model and dates. It skips matching phone numbers already in the CRM and repeated within the same file by default. UploadRow stores normalised data and validation errors; the admin reviews the summary and commits valid rows in a transaction. The server also contains duplicate-resolution actions beyond the current UI defaults.', ['backend/uploads/views.py','backend/uploads/tasks.py','backend/uploads/models.py'])
for i,(h,b) in enumerate([('Upload','CSV / XLSX · ≤10 MB'),('Parse','Normalise and validate'),('Stage / review','Valid · skipped · errors'),('Commit','Leads + audit + batch state')]):
 x=.7+i*3.16;node(s,x,2.85,2.85,h,b,1.23)
 if i<3:line(s,x+2.85,3.47,x+3.14,3.47,arrow=True)
card(s,.7,4.65,5.82,1.36,'Duplicate handling','CRM matches and repeated file phones default to skip.',size=16)
card(s,6.76,4.65,5.82,1.36,'State tracking','PARSING → READY → COMMITTED; failures → FAILED.',accent='blue',size=16)
band(s,'Manual lead entry does not have a unique-phone database constraint; import checks are a separate control.')

s=new('Complaint management workflow','Customer support','CE/reception log the issue; the complaints department handles resolution.',
 'New complaints require a category and a matching subtype, alongside customer and issue details. CE and receptionist users can create and view their own tickets; reception complaints default to walk-in. The complaints department handles updates and notes. Admin can review tickets and complaint analytics. Open, in progress, escalated, resolved and closed are model states; this diagram describes the typical workflow rather than a fully enforced transition graph. The subtype list is an intake taxonomy and must not be represented as evidence of actual vehicle defects.', ['backend/complaints/models.py','backend/complaints/serializers.py','backend/complaints/permissions.py','backend/complaints/test_subtypes.py'])
for i,(h,b) in enumerate([('Open','Category + subtype + priority'),('In progress','Investigate and add notes'),('Resolved','Record resolution notes'),('Closed','Final ticket status')]):
 x=.7+i*3.16;node(s,x,2.85,2.85,h,b,1.27)
 if i<3:line(s,x+2.85,3.48,x+3.13,3.48,arrow=True)
node(s,4.37,4.9,4.56,'Escalated','Escalation state for issues requiring attention',1.05,'blue')
line(s,5.28,4.12,5.28,4.9,'blue',2,arrow=True)
band(s,'ComplaintNote stores discussion history. A related lead is optional; no automatic lead-to-ticket integration is implied.')

s=new('ETBR: define the numbers before comparing them','Analytics','Enquired → Test Drive Completed → Booked → Retailed',
 'ETBR counts distinct enquiries within the caller’s authorised and date-filtered cohort. Enquired counts leads. Test Drive Completed counts non-null completion timestamps. Booked includes both BOOKED and RETAILED sales outcomes. Retailed counts only RETAILED. A booked lead can lack a recorded test drive, so do not assume a strictly decreasing funnel across all four values. These are stage counts within an enquiry cohort, not a historical event-timing funnel or measured company performance.', ['backend/leads/metrics.py','backend/analytics/test_etbr.py','frontend/src/components/etbr-tiles.tsx'])
for i,(letter,h,b) in enumerate([('E','Enquired','Distinct enquiry records'),('T','Test Drive Completed','Completion timestamp exists'),('B','Booked','BOOKED or RETAILED'),('R','Retailed','RETAILED outcome')]):
 x=.7+i*3.16;box(s,x,2.69,2.85,2.91,'white');text(s,x+.2,2.96,2.45,.78,letter,48,'red',True);text(s,x+.2,4.0,2.45,.65,h,20,'ink',True);text(s,x+.2,4.94,2.45,.43,b,14,'muted')
band(s,'Counts only: no fabricated conversion percentages, customer totals or business-impact claims.')

cards('Reporting, performance and caching','Analytics','Comparable metrics require matching permissions, branch and date filters.',[
 ('Reporting views','Admin overview\nPersonal employee analytics\nSales-manager branch reports'),('Database efficiency','select_related on lead ownership\nAggregate counts and next follow-up\nIndexes for common queue filters'),('Cache design','Key includes user, role and location\nQuery parameters included in key\nCache failure falls back to computation')],
 'analytics/cache.py builds a SHA-256 key from endpoint, user identity, role, location and normalised query parameters. This avoids sharing one user’s cached report with another. Only successful responses are cached. CACHE_TTL_SECONDS controls caching; the Render definition sets 50 seconds, while the application default is zero. Cache exceptions produce X-Cache: ERROR and compute the report. The implementation has performance-conscious choices, but no load test or measured speedup is claimed.', ['backend/analytics/cache.py','backend/analytics/views.py','backend/leads/views.py','render.yaml'])

s=new('Employee offboarding without losing work','Operational continuity','Preview impact → choose replacement or queue → apply the change → review held reminders',
 'CE and PS/SO offboarding is a business workflow rather than a simple user delete. The impact preview includes actionable leads, closed work, follow-ups and relevant complaints. A snapshot version detects stale previews and returns a conflict. Within the transaction, active work is routed to eligible recipients or a reassignment queue; the implementation preserves historical attribution and records lifecycle events. Reminders can be held for admin review. The UI label Permanent Delete corresponds to an application lifecycle action, not physical erasure of every related row.', ['backend/accounts/offboarding.py','backend/accounts/views.py','backend/accounts/models.py'])
for i,(h,b) in enumerate([('Impact preview','Counts + snapshot version'),('Choose routing','Eligible replacement / queue'),('Apply lifecycle','Lock rows + preserve history'),('Review reminders','Approve next owner and date')]):
 x=.7+i*3.16;node(s,x,3.04,2.85,h,b,1.27)
 if i<3:line(s,x+2.85,3.67,x+3.14,3.67,arrow=True)
text(s,1.06,5.01,11.15,.63,'Stale preview → HTTP 409 → refresh the impact → review the latest work before retrying',19,'blue',True,True)
band(s,'Business lesson: an employee’s account lifecycle and the customer’s work lifecycle must be handled together.')

s=new('Deployment topology and configuration','Deployment','Repository configuration describes the intended hosting boundaries, not proof of a live rollout.',
 'The frontend uses NEXT_PUBLIC_API_URL and the project has Vercel deployment metadata. render.yaml defines the Python API service, build script, Gunicorn with Uvicorn worker and /api/schema/ health check. DATABASE_URL and Supabase credentials are configured externally. The current service sets CELERY_TASK_ALWAYS_EAGER=true and does not define a separate worker or beat service. Therefore imports can run inline and automatic due reminders need independent scheduling. No production system or database was queried for this presentation.', ['render.yaml','backend/build.sh','backend/config/settings.py','frontend/src/lib/crm.ts'])
node(s,.75,2.79,3.1,'Frontend host','Vercel / Next.js configuration',1.17,'blue')
node(s,4.55,2.79,4.18,'API host','Render · Gunicorn / Uvicorn',1.17)
node(s,9.42,2.79,3.14,'Database','DATABASE_URL configured externally',1.17,'green')
line(s,3.85,3.37,4.55,3.37,'blue',2,arrow=True);line(s,8.73,3.37,9.42,3.37,'green',2,arrow=True)
node(s,.75,4.9,3.8,'File storage','Supabase when configured',1.1,'blue')
node(s,4.94,4.9,3.35,'Analytics cache','Redis / memory fallback',1.1,'blue')
node(s,8.69,4.9,3.87,'Future task services','Worker + scheduler deployment',1.1,'muted')
for x in [2.65,6.61,10.62]:line(s,6.64,3.96,x,4.9,'muted',1.5,arrow=True,dash=True)
band(s,'Deployment handover: migrations, secrets, allowed origins, health checks and role-by-role acceptance.')

cards('Testing and current verification','Quality assurance','Verification performed while preparing this presentation; results recorded in Verification.md.',[
 ('102 backend tests passed','Lead ownership and handoff\nImports, complaints and offboarding\nETBR, activities and intake rules'),('Frontend static checks','TypeScript: npm run typecheck\nESLint: npm run lint\nBoth completed successfully'),('Validation boundaries','Isolated SQLite test database\nNo live deployment acceptance test\nNo load test or production metrics')],
 'The current run passed 102 backend tests in 98.912 seconds with no Django system-check issues. The tests ran against an isolated test database with eager tasks and caching disabled. They test workflow and permission behaviour, including recent intake and ETBR additions. These checks were run for presentation preparation on 11 September 2026 and should not be described as your original internship-week test execution. Frontend static checks confirm type and lint consistency but do not prove that every browser flow or live infrastructure service is working.', ['backend/leads/tests.py','backend/leads/test_intake.py','backend/analytics/test_etbr.py','backend/complaints/test_subtypes.py','backend/accounts/tests.py','Verification.md'],take='11 September 2026: 102 backend tests passed; TypeScript and ESLint checks passed.')

table('Technical challenges and responses','Engineering learning','Code-backed design issues; personalise the experience with incidents you actually handled.',[
 ['Challenge','Response visible in code','Learning to discuss'],
 ['CE and PS share one customer','Separate owner fields and qualification context','Model responsibility before building forms'],
 ['Repeated or invalid intake data','Validation, normalisation and staged import','Validate before committing customer records'],
 ['Users leave with open work','Impact preview, reassignment and history','Handle business continuity explicitly'],
 ['Reports differ across users','Authorised cohorts and scoped cache keys','Define metric scope before comparing totals'],
 ['Multiple clicks / concurrent updates','Transactions, locks and idempotent completion','Make repeat actions safe where required'],
],[3.15,4.8,4.13],
 'These challenges are supported by implementation choices, but the code cannot establish which difficulties you personally experienced. Replace one or two rows with real examples: the trigger, what you saw, how you diagnosed it, what you changed and how you verified the result. Avoid claiming you independently solved a team problem unless that is accurate. Useful evidence includes the commit, test case and mentor review associated with the change.', ['backend/leads/views.py','backend/uploads/tasks.py','backend/accounts/offboarding.py','backend/analytics/cache.py'],15)

cards('My internship experience and learning','Reflection','Draft reflection themes: personalise these with your own examples before presenting.',[
 ('Understanding business work','Explain an actual showroom observation\nDescribe how it affected a requirement\nMention one useful mentor discussion'),('Growing as a developer','Explain one module you contributed to\nShow a validation or workflow you understood\nDescribe one debugging experience'),('Working professionally','Describe how you asked for feedback\nExplain how you documented progress\nAdd one change in your working habits')],
 'Suggested first-person wording, use only if accurate: “At the beginning, I focused on the screens. As I understood the workflow, I started asking who owns each enquiry and what action should happen next.” Another option: “R&D helped me turn a general request into specific fields, permissions and validation rules.” Add a real event, the people involved by role and the resulting lesson. The presenter has not supplied personal anecdotes, so this slide intentionally uses prompts rather than invented experiences.',take='Personalise: one real observation + one technical contribution + one feedback example.')

cards('Project outcomes and remaining work','Evaluation','Describe implemented capabilities separately from measured business impact.',[
 ('Available in the repository','Role-specific CRM workflows\nSales and complaint histories\nImports, analytics and lifecycle controls'),('Evidence available','Implementation files and migrations\nAutomated workflow checks\nTechnical diagrams and demo scenarios'),('Next steps','Confirm user acceptance and rollout scope\nDeploy and verify reminder scheduling\nMeasure usage and business results')],
 'The defensible outcome is a working implementation represented by code and tests, not a claimed percentage improvement in company sales or employee productivity. Identify the portion you personally delivered. Remaining work includes live acceptance across roles, deployment verification, reminder scheduling and measurement of actual operational outcomes. Other potential extensions include narrower production origins, export and reporting improvements, validated source integrations and accessibility/browser review. Present them as future work, not completed features.')

cards('Demonstration plan','Project demonstration','Use fictional records in an isolated demo environment.',[
 ('Scenario 1: enquiry to sale','Capture and assign a lead\nCE qualifies and selects PS/SO\nPS/SO records booking and retail'),('Scenario 2: customer issue','CE or reception creates complaint\nComplaints team adds notes\nRecord resolution and review history'),('Scenario 3: management','Show the matching ETBR scope\nReview branch visibility\nExplain offboarding impact preview')],
 'Prepare a fictional customer record and use separate browser profiles for different roles because login cookies are shared between tabs in one profile. Start with configured branch names, active staff, model and source values. After each handoff, refresh the receiving queue and show the same lead identifier. Do not read real customer details aloud. If a live environment is unavailable, use the editable workflow diagrams and narrate the expected action and evidence rather than claiming a live demonstration occurred.',take='Demo proof: saved record + next owner’s queue + history + matching report filters.')

s=new('Thank you','Closing','Questions and discussion',
 'Close by connecting the internship to your learning: understanding dealership operations, translating responsibilities into software rules, and checking those rules with evidence. Thank your actual company mentor, college guide and colleagues by name only after confirming the details. Invite questions about the workflow, architecture, database and the specific modules you contributed to. The following appendix supports a more detailed technical viva.',dark=True)
text(s,.73,2.98,11.87,1.22,'From showroom processes\nto software workflows',39,'white',True)
text(s,.76,5.15,11.44,.68,'Nippon Toyota • Kalamassery\n[Your name]  |  [College]  |  [Mentor]',19,'line')

# Technical appendix.
table('API reference for the viva','Technical appendix','Selected implemented endpoints; consult /api/docs/ for the complete schema.',[
 ['Method','Endpoint','Purpose'],
 ['POST','/api/auth/login/','Authenticate and set cookies'],
 ['GET / POST','/api/leads/','List visible leads / create allowed intake'],
 ['PATCH','/api/leads/{id}/so-update/','Dedicated call and sales workflow update'],
 ['POST','/api/leads/{id}/assign/','Allocate to CE'],
 ['POST','/api/leads/{id}/complete-test-drive/','Record completion once'],
 ['POST','/api/uploads/{id}/commit/','Commit reviewed import rows'],
 ['GET','/api/analytics/sales-manager/','Branch-scoped analytics'],
],[1.65,6.16,4.27],
 'Resource routes are registered through Django REST Framework routers, while auth and analytics include explicit paths. Endpoint naming retains earlier terminology: so-update handles the dedicated business update route even though CE and PS/SO are separate user roles. Explain request validation, authentication and scope before discussing the result. Status codes include 400 for validation errors, 401 for missing/invalid authentication, 403 for forbidden actions, 404 for objects outside a filtered queryset and 409 for stale offboarding snapshots.', ['backend/config/urls.py','backend/leads/urls.py','backend/accounts/urls.py','backend/uploads/urls.py','backend/analytics/urls.py'],14)

table('Additional entities and persistence choices','Technical appendix','Operational records around the core Lead model.',[
 ['Entity','Key relationships / fields','Why it exists'],
 ['UploadBatch → UploadRow','One batch has many staged rows','Review validation and duplicate results'],
 ['Complaint → ComplaintNote','One ticket has many notes','Keep investigation discussion'],
 ['Notification','User; optional Lead; read_at','Persist assignment and follow-up messages'],
 ['UserLifecycleEvent','User, actor, reason and JSON summary','Retain disable/enable/delete history'],
 ['SystemConfig','JSON lists and updated_at','Manage sources, models and activity hierarchy'],
 ['LeadAudit','Lead, actor, before and after JSON','Explain who changed customer workflow data'],
],[3.13,4.7,4.25],
 'Django migrations evolve the schema. Optional or nullable fields support existing records while newer serializers enforce requirements for new submissions. UserLifecycleEvent protects the target user relation; several author references use PROTECT so historical records retain their identity. Other optional relationships use SET_NULL. Distinguish a historical audit record from a backup: audits do not replace database backup, recovery and retention procedures. Those operational procedures were not verified.', ['backend/uploads/models.py','backend/complaints/models.py','backend/notifications/models.py','backend/accounts/models.py','backend/leads/models.py'],15)

table('Technical viva: likely questions','Technical appendix','Short answers to rehearse, with implementation-specific distinctions.',[
 ['Question','Answer to develop'],
 ['Why separate frontend and API?','Role-specific UI stays separate from shared validation and authorisation.'],
 ['Why a relational database?','Ownership, calls, follow-ups and tickets have explicit relationships and transactional updates.'],
 ['Why two lead owners?','The CE relationship remains while PS/SO takes responsibility for the sales stage.'],
 ['How do you prevent duplicates?','Import parsing skips repeated phones; manual entry has no unique-phone constraint.'],
 ['What makes a repeated action safe?','Test-drive completion locks the lead and writes its timestamp/audit only once.'],
 ['Are reminders live?','Tasks and a schedule exist; the checked-in Render service needs separate scheduling verification.'],
],[3.9,8.18],
 'Use these as starting points and answer in your own words. When asked about framework choice, connect the choice to forms, APIs, authorisation and relational data rather than generic popularity. When asked about security, explain the exact controls and acknowledge review limits. When asked what you built, state only your verified contribution. When asked about database deployment, distinguish SQLite defaults, PostgreSQL support and the actual provider, which has not been inspected.', ['backend/config/settings.py','backend/leads/views.py','backend/uploads/tasks.py','render.yaml'],15)

table('References and photo credits','Sources','Company sources accessed 11 September 2026; technical content reviewed from the local repository.',[
 ['Source','What it supports'],
 ['Nippon Toyota official About Us / CMS','Mission; Dealer Principal and Directors'],
 ['Official Kalamassery location page','Legal entity, showroom/service centre and address'],
 ['Company LinkedIn; official U Trust page','Published establishment date; used-car offering'],
 ['Ranjithsiji / Wikimedia Commons (2012)','Kalamassery Toyota Office.JPG; CC BY-SA 3.0; unmodified photo'],
 ['Sreejith Sivadasan / Flickr (2011)','Nippon Toyota showroom at South Kalamassery; all rights reserved'],
 ['Local repository and Verification.md','Architecture, models, workflows, configuration and checks'],
],[5.35,6.73],
 'Full URLs and photo licensing details are supplied in Sources_and_Customisation.md and embedded in the notes. The Wikimedia photograph is licensed CC BY-SA 3.0: credit Ranjithsiji, link the source and licence, and retain the licence for any adapted version. It is placed unmodified in this presentation. The Flickr photograph retains its photographer’s copyright; public availability is not an open licence. Seek the owner’s permission or replace it with a personal/licensed image before public redistribution. Neither photograph is a current office-interior record.', list(SOURCES.values()),15)

for s,d in zip(prs.slides,slides):
 n=d['notes']+'\n\nSources / implementation evidence:\n'+'\n'.join(d['refs']) if d['refs'] else d['notes']
 s.notes_slide.notes_text_frame.text=n
prs.save(P/'Nippon_Toyota_Internship_Presentation.pptx')
(P/'source/slide_content.json').write_text(json.dumps(slides,indent=2,ensure_ascii=False))

notes=['# Nippon Toyota internship: presenter notes','', 'Eight-week draft timeline. Confirm personal details, contribution, dates and experience before presenting.','']
for d in slides:
 notes += [f"## {d['number']:02d}. {d['title']}",'',d['notes'],'']
 if d['refs']:notes += ['Evidence: '+ '; '.join(d['refs']),'']
(P/'Presenter_Notes.md').write_text('\n'.join(notes))
pdfmetrics.registerFont(TTFont('Lato','/usr/share/fonts/truetype/lato/Lato-Regular.ttf'))
pdfmetrics.registerFont(TTFont('LatoBold','/usr/share/fonts/truetype/lato/Lato-Bold.ttf'))
styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='BodyL',fontName='Lato',fontSize=11,leading=16,spaceAfter=12));styles.add(ParagraphStyle(name='HeadL',fontName='LatoBold',fontSize=21,leading=27,textColor=HexColor('#C8202F'),spaceAfter=20));styles.add(ParagraphStyle(name='SmallL',fontName='Lato',fontSize=9,leading=13,textColor=HexColor('#59636A'),spaceAfter=9))
story=[]
for d in slides:
 story.extend([Paragraph(f"{d['number']:02d} / {escape(d['section'].upper())}",styles['SmallL']),Paragraph(escape(d['title']),styles['HeadL']),Paragraph(escape(d['notes']),styles['BodyL']),Spacer(1,10),Paragraph('Visible slide content',styles['SmallL']),Paragraph(escape(' • '.join(d['text'][3:-2])).replace('\n','<br/>'),styles['SmallL'])])
 if d['refs']:story.extend([Spacer(1,12),Paragraph('Sources / implementation evidence',styles['SmallL']),Paragraph('<br/>'.join(escape(v) for v in d['refs']),styles['SmallL'])])
 story.append(PageBreak())
SimpleDocTemplate(str(P/'Presenter_Notes.pdf'),rightMargin=45,leftMargin=45,topMargin=40,bottomMargin=40).build(story)
print(f'Created {len(slides)} editable slides, embedded notes and presenter guide.')
