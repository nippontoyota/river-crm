from pathlib import Path
from collections import Counter
import re,json,subprocess,xml.etree.ElementTree as E
from pptx import Presentation
P=Path(__file__).resolve().parents[1]
s=json.loads((P/'source/slide_content.json').read_text()); deck=Presentation(P/'Incheon_Mobility_Operations_Demo.pptx')
assert len(s)==len(deck.slides)==30
for i,sl in enumerate(deck.slides):
 assert s[i]['notes'] in sl.notes_slide.notes_text_frame.text
 for sh in sl.shapes:
  assert sh.left>=0 and sh.top>=0
  assert sh.left+sh.width<=deck.slide_width+10000 and sh.top+sh.height<=deck.slide_height+10000
pdftext=subprocess.check_output(['pdftotext',str(P/'Incheon_Mobility_Operations_Demo.pdf'),'-']).decode()
pages=pdftext.split('\f');assert len(pages)-1==30
norm=lambda t:re.findall('[a-z0-9]+',t.casefold().replace('-', ''))
issues=[]
for i,(d,page) in enumerate(zip(s,pages),1):
 expected=Counter(norm(' '.join(d['visible_text'])))
 actual=Counter(norm(page))
 missing=expected-actual
 if missing:issues.append((i,dict(missing)))
assert not issues,issues
bbox=subprocess.check_output(['pdftotext','-bbox-layout',str(P/'Incheon_Mobility_Operations_Demo.pdf'),'-'])
r=E.fromstring(bbox);ns={'h':'http://www.w3.org/1999/xhtml'}
for i,page in enumerate(r.findall('.//h:page',ns),1):
 for word in page.findall('.//h:word',ns):
  assert 0<=float(word.attrib['xMin']) and float(word.attrib['xMax'])<=float(page.attrib['width']), (i,word.text)
  assert 0<=float(word.attrib['yMin']) and float(word.attrib['yMax'])<=float(page.attrib['height']), (i,word.text)
print('PASS: 30 slides; all embedded notes; shape bounds; all slide-body words retained in PDF; no PDF text outside page bounds.')
